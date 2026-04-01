from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os


def generate_template(scope, stakeholders, cert, task, comm):

    # ----------------------------
    # FUNCTIONS
    # ----------------------------
    def add_page_border(section):
        sectPr = section._sectPr
        pgBorders = OxmlElement('w:pgBorders')
        pgBorders.set(qn('w:offsetFrom'), 'page')

        for side in ['top', 'left', 'bottom', 'right']:
            border = OxmlElement(f'w:{side}')
            border.set(qn('w:val'), 'single')
            border.set(qn('w:sz'), '12')
            border.set(qn('w:space'), '24')
            border.set(qn('w:color'), '000000')
            pgBorders.append(border)

        sectPr.append(pgBorders)

    def remove_page_border(section):
        sectPr = section._sectPr
        for el in sectPr.findall(qn('w:pgBorders')):
            sectPr.remove(el)

    def heading(doc, text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)

    # ----------------------------
    # CREATE DOCUMENT
    # ----------------------------
    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # ----------------------------
    # SECTION 1 → COVER PAGE
    # ----------------------------
    section1 = doc.sections[0]
    section1.top_margin = Inches(1.5)
    section1.bottom_margin = Inches(1.5)
    section1.left_margin = Inches(1.5)
    section1.right_margin = Inches(1.5)

    add_page_border(section1)

    doc.add_paragraph("File No. ________________________")

    p = doc.add_paragraph()
    run = p.add_run("\nTASK DIRECTIVE\n")
    run.bold = True
    run.font.size = Pt(20)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("________ /2026").alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Issue No."
    table.cell(0, 1).text = "Date of Issue:"

    doc.add_paragraph("\nPROJECT NAME: __________________________________________")

    p = doc.add_paragraph("\n\n[ LOGO HERE ]")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("ABC Organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Address of organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Organization details").alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ----------------------------
    # SECTION 2 → MAIN CONTENT
    # ----------------------------
    section2 = doc.add_section(WD_SECTION.NEW_PAGE)
    section2.header.is_linked_to_previous = False
    section2.footer.is_linked_to_previous = False

    remove_page_border(section2)

    # ----------------------------
    # FULL TEMPLATE CONTENT
    # ----------------------------

    heading(doc, "1. Introduction [Automate]")
    doc.add_paragraph("__________________________________________________")

    heading(doc, "2. Reference [Default]")
    doc.add_paragraph("__________________________________________________")

    heading(doc, "3. Basis Of Task Directive [Default]")
    doc.add_paragraph("__________________________________________________")

    # ✅ Injected
    heading(doc, "4. Scope Of Task Directive [Drop down menu]")
    doc.add_paragraph(scope)

    # ✅ Injected
    heading(doc, "5. Stakeholders [Automate + Manual]")
    doc.add_paragraph(f"The following are the major stakeholders: {stakeholders}")

    table = doc.add_table(rows=4, cols=4)
    table.style = 'Table Grid'
    headers = ["Sl No.", "Organisation", "Role", "Activities"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    # ✅ Injected
    heading(doc, "6. Certification Work Breakdown [Drop down menu]")
    doc.add_paragraph(cert)

    # ✅ Injected
    heading(doc, "7. Task Allocation [Default]")
    doc.add_paragraph(task)

    heading(doc, "7.1 Coordinating Directorate [Default]")
    doc.add_paragraph("______________________________________________")

    heading(doc, "7.2 Single Point of Contact (SPoC) [Default]")
    doc.add_paragraph("______________________________________________")

    heading(doc, "7.3 Certification Task Allocation")

    table = doc.add_table(rows=4, cols=4)
    table.style = 'Table Grid'
    headers = ["Sl No.", "Certification Activity", "Certification Work Centre", "Responsible Head"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    heading(doc, "7.4 Issue of Clearance [Default]")
    doc.add_paragraph("a. ___\nb. ___\nc. ___\nd. ___\ne. ___\nf. ___\ng. ___")

    heading(doc, "8. SCRB And TARB [Default]")
    doc.add_paragraph("______________________________________________")

    # ✅ Injected
    heading(doc, "9. Communication [Default]")
    doc.add_paragraph(comm)

    heading(doc, "10. Certification Progress Review [Default]")
    doc.add_paragraph("______________________________________________\n\n(__________)")

    heading(doc, "11. Distribution List")
    doc.add_paragraph("11.1 External Organization\n1. ___\n2. ___\n3. ___")
    doc.add_paragraph("11.2 Internal Distribution")

    # ----------------------------
    # ANNEXURE 1
    # ----------------------------
    doc.add_page_break()

    heading(doc, "Annexure-1")
    doc.add_paragraph("Work Assignment List of LRUs [Automate + Manual]")

    table = doc.add_table(rows=4, cols=6)
    table.style = 'Table Grid'
    headers = ["Sl No", "ABC", "Abc2", "Abc3", "Abc4", "Abc5"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    # ----------------------------
    # ANNEXURE 2
    # ----------------------------
    doc.add_page_break()

    heading(doc, "Annexure-2")
    doc.add_paragraph("Contact details of dealing officers and RDs [Manual]")

    table = doc.add_table(rows=3, cols=5)
    table.style = 'Table Grid'
    headers = ["Sl no.", "Abc1", "Abc2", "Abc3", "Abc4"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h

    # ----------------------------
    # ANNEXURE 3
    # ----------------------------
    doc.add_page_break()

    heading(doc, "Annexure-3")
    doc.add_paragraph("Product Break Down Structure [Automate]")

    # ----------------------------
    # SAVE
    # ----------------------------
    output_path = os.path.join(os.path.dirname(__file__), "generated_template.docx")
    doc.save(output_path)

    return output_path
