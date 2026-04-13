import sys
import re
import fitz  # PyMuPDF
import tempfile 
import os       

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout,
    QHBoxLayout, QPushButton, QTextEdit, QFileDialog,
    QLabel, QSplitter, QMessageBox, QScrollArea, QProgressBar,
    QCheckBox, QGroupBox, QDialog, QRubberBand, QMenu, QLineEdit,
    QGridLayout, QListWidget, QListWidgetItem
)
from PyQt6.QtGui import QPixmap, QImage, QFont, QTextDocument, QTextCursor, QDesktopServices 
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QEvent, QRect, QPoint, QSize, QUrl 
from PyQt6.QtPrintSupport import QPrinter

# --- WORD TEMPLATE IMPORTS ---
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ==========================================
# --- CUSTOM UI: INLINE DROPDOWN WITH [+] ---
# ==========================================
class DropdownPopup(QDialog):
    def __init__(self, parent_btn, options):
        super().__init__(parent_btn, Qt.WindowType.Popup | Qt.WindowType.FramelessWindowHint)
        self.parent_btn = parent_btn
        self.setStyleSheet("""
            QDialog { background: #FFFFFF; border: 1px solid #ced4da; border-radius: 4px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); }
            QLabel { color: #6B7280; font-size: 11px; font-weight: bold; }
            QListWidget { background: transparent; border: none; color: #111827; outline: none; font-size: 13px; }
            QListWidget::item { padding: 6px; border-radius: 4px; }
            QListWidget::item:hover { background: #F3F4F6; }
            QLineEdit { background: #FFFFFF; color: #111827; border: 1px solid #2b579a; border-radius: 4px; padding: 4px; }
        """)
        
        layout = QVBoxLayout(self)
        layout.setContentsMargins(6, 6, 6, 6)
        layout.setSpacing(6)

        # Header: Title + Add Button
        header_layout = QHBoxLayout()
        header_layout.addWidget(QLabel("SELECT OPTIONS"))
        header_layout.addStretch()
        
        self.add_btn = QPushButton("+")
        self.add_btn.setFixedSize(22, 22)
        self.add_btn.setStyleSheet("""
            QPushButton { background: #F3F4F6; color: #111827; border: 1px solid #D1D5DB; border-radius: 4px; font-weight: bold; font-size: 14px; }
            QPushButton:hover { background: #E5E7EB; }
        """)
        self.add_btn.clicked.connect(self.show_new_input)
        header_layout.addWidget(self.add_btn)
        layout.addLayout(header_layout)

        # Hidden Input for new items
        self.new_input = QLineEdit()
        self.new_input.setPlaceholderText("Type & press Enter...")
        self.new_input.hide()
        self.new_input.returnPressed.connect(self.save_new_option)
        layout.addWidget(self.new_input)

        # List of Checkboxes
        self.list_widget = QListWidget()
        for opt in options:
            if opt != "Enter Other...":
                self.add_item(opt)
        self.list_widget.itemChanged.connect(self.parent_btn.update_display)
        layout.addWidget(self.list_widget)

    def add_item(self, text, checked=False):
        item = QListWidgetItem(text)
        item.setFlags(item.flags() | Qt.ItemFlag.ItemIsUserCheckable)
        item.setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)
        self.list_widget.addItem(item)

    def show_new_input(self):
        self.new_input.show()
        self.new_input.setFocus()

    def save_new_option(self):
        text = self.new_input.text().strip()
        if text:
            existing = [self.list_widget.item(i).text() for i in range(self.list_widget.count())]
            if text not in existing:
                self.add_item(text, checked=True)
            else:
                for i in range(self.list_widget.count()):
                    if self.list_widget.item(i).text() == text:
                        self.list_widget.item(i).setCheckState(Qt.CheckState.Checked)
        self.new_input.clear()
        self.new_input.hide()
        self.parent_btn.update_display()

class PopupMultiSelect(QPushButton):
    def __init__(self, options, parent=None):
        super().__init__("Select options...", parent)
        self.setStyleSheet("""
            QPushButton {
                text-align: left; padding: 8px; background: #FFFFFF; 
                border: 1px solid #ced4da; border-radius: 4px; color: #111827; font-size: 13px;
            }
            QPushButton:hover { border: 1px solid #2b579a; }
        """)
        self.options = options
        self.popup = None

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            if not self.popup:
                self.popup = DropdownPopup(self, self.options)
            
            pos = self.mapToGlobal(self.rect().bottomLeft())
            self.popup.setFixedWidth(self.width())
            self.popup.move(pos.x(), pos.y() + 2)
            self.popup.show()

    def update_display(self):
        checked = self.checkedItems()
        if checked:
            self.setText(", ".join(checked))
        else:
            self.setText("Select options...")

    def checkedItems(self):
        if not self.popup: return []
        items = []
        for i in range(self.popup.list_widget.count()):
            item = self.popup.list_widget.item(i)
            if item.checkState() == Qt.CheckState.Checked:
                items.append(item.text())
        return items


# ==========================================
# --- PURE PYTHON SUMMARIZER ---
# ==========================================
def simple_summarize(text, num_sentences=4):
    if not text: return ""
    stop_words = {"the", "is", "in", "and", "to", "of", "a", "for", "on", "with", "as", "by", "this", "that", "it", "are", "be", "or", "an", "at", "from", "which", "will"}
    sentences = re.split(r'(?<=[.!?]) +|\n+', text)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 5]
    if len(sentences) <= num_sentences: return " ".join(sentences)
        
    words = re.findall(r'\b\w+\b', text.lower())
    freq = {}
    for w in words:
        if w not in stop_words and not w.isnumeric(): freq[w] = freq.get(w, 0) + 1
            
    scores = {}
    for i, s in enumerate(sentences):
        score = 0
        s_words = re.findall(r'\b\w+\b', s.lower())
        for w in s_words:
            if w in freq: score += freq[w]
        scores[i] = score / max(len(s_words), 1)
        
    top_indices = sorted(sorted(scores, key=scores.get, reverse=True)[:num_sentences])
    return " ".join([sentences[i] for i in top_indices])

# ==========================================
# --- SMART TEXT EDITOR (SUPPORTS IMAGES)---
# ==========================================
class NotesEditor(QTextEdit):
    def canInsertFromMimeData(self, source):
        if source.hasImage() or source.hasText(): return True
        return super().canInsertFromMimeData(source)

    def insertFromMimeData(self, source):
        if source.hasImage():
            clipboard = QApplication.clipboard()
            image = clipboard.image() 
            if not image.isNull():
                cursor = self.textCursor()
                if image.width() > 500:
                    image = image.scaledToWidth(500, Qt.TransformationMode.SmoothTransformation)
                cursor.insertImage(image)
                self.insertPlainText("\n")
                return 
        if source.hasText():
            self.insertPlainText(source.text())
            return
        super().insertFromMimeData(source)

# ==========================================
# --- PURE CODE ANALYSIS WORKER (SMART EXTRACTION) ---
# ==========================================
class AnalysisWorker(QThread):
    finished = pyqtSignal(dict) 
    error = pyqtSignal(str)

    def __init__(self, text, topics):
        super().__init__()
        self.text = text
        self.topics = topics

    def run(self):
        try:
            results = {
                "project_name": "",
                "introduction": "",
                "found_topics": []
            }

            explicit_match = re.search(r'(?:Project\s*Name|Project\s*Title|Title|Subject)[\s:]*(.+)', self.text, re.IGNORECASE)
            if explicit_match and not explicit_match.group(1).strip().startswith("___"):
                results["project_name"] = explicit_match.group(1).strip()[:100]
            
            if not results["project_name"]:
                semantic_match = re.search(r'(?:project|report|proposal) (?:aims to|focuses on|proposes|is to|investigates) ([^\.]+)', self.text, re.IGNORECASE)
                if semantic_match:
                    results["project_name"] = semantic_match.group(1).strip().title()[:100]

            if not results["project_name"]:
                lines = [line.strip() for line in self.text.split('\n') if line.strip()]
                for line in lines[:10]:
                    if re.search(r'(Document Analysis Report|Summary|Page|Date|Author)', line, re.IGNORECASE): continue
                    if 10 < len(line) < 100:
                        results["project_name"] = line
                        break

            results["introduction"] = simple_summarize(self.text, num_sentences=4)

            for topic in self.topics:
                if topic.lower() in self.text.lower():
                    results["found_topics"].append(topic)

            if results["introduction"] and "Introduction" not in results["found_topics"]:
                results["found_topics"].append("Introduction")

            self.finished.emit(results)
        except Exception as e:
            self.error.emit(str(e))

# ==========================================
# --- INTERACTIVE PDF SNIPPING & PANNING TOOL ---
# ==========================================
class PdfPageLabel(QLabel):
    text_extracted = pyqtSignal(str)
    image_extracted = pyqtSignal(QPixmap)
    summary_requested = pyqtSignal(str) 

    def __init__(self, page, scroll_area, parent=None):
        super().__init__(parent)
        self.page = page  
        self.scroll_area = scroll_area 
        self.rubber_band = QRubberBand(QRubberBand.Shape.Rectangle, self)
        self.origin = QPoint()
        self.pan_start_pos = None

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.origin = event.pos()
            self.rubber_band.setGeometry(QRect(self.origin, QSize()))
            self.rubber_band.show()
        elif event.button() in [Qt.MouseButton.RightButton, Qt.MouseButton.MiddleButton]:
            self.pan_start_pos = event.globalPosition().toPoint()
            self.setCursor(Qt.CursorShape.ClosedHandCursor)

    def mouseMoveEvent(self, event):
        if not self.origin.isNull() and (event.buttons() & Qt.MouseButton.LeftButton):
            self.rubber_band.setGeometry(QRect(self.origin, event.pos()).normalized())
            
        elif self.pan_start_pos is not None and (event.buttons() & (Qt.MouseButton.RightButton | Qt.MouseButton.MiddleButton)):
            current_pos = event.globalPosition().toPoint()
            delta = current_pos - self.pan_start_pos
            
            h_bar = self.scroll_area.horizontalScrollBar()
            v_bar = self.scroll_area.verticalScrollBar()
            
            h_bar.setValue(h_bar.value() - delta.x())
            v_bar.setValue(v_bar.value() - delta.y())
            
            self.pan_start_pos = current_pos

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            rect = self.rubber_band.geometry()
            if rect.width() > 10 and rect.height() > 10:
                self.show_context_menu(event.pos(), rect)
            else:
                self.rubber_band.hide()
            self.origin = QPoint()
            
        elif event.button() in [Qt.MouseButton.RightButton, Qt.MouseButton.MiddleButton]:
            self.unsetCursor()
            self.pan_start_pos = None
            
    def get_text_from_rect(self, rect):
        if self.pixmap() and self.page:
            scale_x = self.page.rect.width / self.pixmap().width()
            scale_y = self.page.rect.height / self.pixmap().height()
            pdf_rect = fitz.Rect(rect.left() * scale_x, rect.top() * scale_y, rect.right() * scale_x, rect.bottom() * scale_y)
            pdf_rect = pdf_rect + (-3, -3, 3, 3)
            return self.page.get_textbox(pdf_rect).strip()
        return ""

    def show_context_menu(self, pos, rect):
        self.rubber_band.hide() 
        menu = QMenu(self)
        menu.setStyleSheet("""
            QMenu { background-color: white; border: 1px solid #ced4da; border-radius: 4px; padding: 4px; font-size: 13px; }
            QMenu::item { padding: 6px 25px 6px 20px; background-color: transparent; }
            QMenu::item:selected { background-color: #e2e6ea; color: black; border-radius: 3px; }
        """)

        copy_img_action = menu.addAction("🖼️ Extract as Image")
        copy_text_action = menu.addAction("📝 Extract Text to Notes")
        menu.addSeparator()
        summarize_action = menu.addAction("✨ Summarize Selection (Set as Intro)")

        action = menu.exec(self.mapToGlobal(pos))
        
        if action == copy_img_action:
            if self.pixmap(): self.image_extracted.emit(self.pixmap().copy(rect)) 
        elif action == copy_text_action:
            text = self.get_text_from_rect(rect)
            if text: self.text_extracted.emit(text) 
        elif action == summarize_action:
            text = self.get_text_from_rect(rect)
            if text: self.summary_requested.emit(text)

# ==========================================
# --- MAIN APPLICATION ---
# ==========================================
class OfflineApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Offline Document Assistant - Fast Analyzer")
        self.resize(1450, 950) 
        
        self.file_path = ""
        self.doc = None
        self.page_data = [] 
        self.current_zoom = 600 
        self.is_maximized = False
        
        self.intro_text_data = "" 
        self.combos = {}
        
        self.init_ui()

    def init_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        layout = QVBoxLayout(main_widget)
        self.splitter = QSplitter(Qt.Orientation.Horizontal)

        # --- LEFT SIDE: Document Viewer ---
        left_container = QWidget()
        left_layout = QVBoxLayout(left_container)
        left_controls_layout = QHBoxLayout()
        
        self.upload_btn = QPushButton(" 📂 Upload PDF Document")
        self.upload_btn.setStyleSheet("height: 40px; font-weight: bold; background-color: #f8f9fa; padding: 0 15px;")
        self.upload_btn.clicked.connect(self.upload_file)
        
        self.zoom_out_btn = QPushButton("➖")
        self.zoom_in_btn = QPushButton("➕")
        self.maximize_btn = QPushButton("🗖") 
        for btn in [self.zoom_out_btn, self.zoom_in_btn, self.maximize_btn]:
            btn.setFixedSize(40, 40)
            btn.setStyleSheet("font-size: 16px; font-weight: bold; background-color: #ffffff; border: 1px solid #ced4da; border-radius: 4px;")
            
        self.zoom_out_btn.clicked.connect(lambda: self.zoom_out(150))
        self.zoom_in_btn.clicked.connect(lambda: self.zoom_in(150))
        self.maximize_btn.clicked.connect(self.toggle_maximize)

        left_controls_layout.addWidget(self.upload_btn)
        left_controls_layout.addStretch()
        left_controls_layout.addWidget(QLabel("Zoom:"))
        left_controls_layout.addWidget(self.zoom_out_btn)
        left_controls_layout.addWidget(self.zoom_in_btn)
        left_controls_layout.addWidget(self.maximize_btn)

        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setStyleSheet("background-color: #525659;") 
        self.page_container = QWidget()
        self.page_layout = QVBoxLayout(self.page_container)
        self.page_layout.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignHCenter) 
        self.scroll_area.setWidget(self.page_container)
        self.scroll_area.viewport().installEventFilter(self)

        self.text_preview = QTextEdit()
        self.text_preview.setReadOnly(True)
        self.text_preview.hide()

        left_layout.addLayout(left_controls_layout)
        tip_label = QLabel("💡 Tip: Click and drag on the PDF to extract text, flowcharts, or generate your Introduction Summary!")
        tip_label.setStyleSheet("color: #0056b3; font-style: italic;")
        left_layout.addWidget(tip_label)
        left_layout.addWidget(self.scroll_area)
        left_layout.addWidget(self.text_preview)

        # --- RIGHT SIDE: Tools & Dashboards ---
        self.right_container = QWidget() 
        right_layout = QVBoxLayout(self.right_container)

        header_layout = QHBoxLayout()

        # Checkbox Grid
        self.status_group = QGroupBox("Task Directive Extraction Checklist")
        self.status_group.setMaximumHeight(190) 
        
        status_main_layout = QVBoxLayout()
        status_main_layout.setContentsMargins(10, 15, 10, 10) 
        status_main_layout.setSpacing(5) 
        
        top_status = QHBoxLayout()
        self.check_upload = QCheckBox("File Uploaded")
        self.check_upload.setEnabled(False)
        self.status_label = QLabel("Ready")
        self.status_label.setStyleSheet("color: gray; font-style: italic;")
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0) 
        self.progress_bar.hide()
        
        top_status.addWidget(self.check_upload)
        top_status.addStretch()
        top_status.addWidget(self.status_label)
        top_status.addWidget(self.progress_bar) 
        status_main_layout.addLayout(top_status)
        
        self.topics = [
            "Introduction", "Reference", "Basis Of Task Directive",
            "Scope Of Task Directive", "Stakeholders", "Certification Work Breakdown",
            "Task Allocation", "Coordinating Directorate", "Single Point of Contact",
            "Certification Task Allocation", "Issue of Clearance", "SCRB And TARB",
            "Communication", "Certification Progress Review", "Distribution List"
        ]
        self.topic_checkboxes = {}
        
        topics_scroll = QScrollArea()
        topics_scroll.setWidgetResizable(True)
        topics_scroll.setFixedHeight(120) 
        topics_scroll.setStyleSheet("QScrollArea { border: 1px solid #ced4da; border-radius: 4px; background-color: #f8f9fa; }")
        
        topics_widget = QWidget()
        topics_layout = QGridLayout(topics_widget) 
        topics_layout.setContentsMargins(5, 5, 5, 5)
        
        row, col = 0, 0
        for topic in self.topics:
            cb = QCheckBox(topic)
            cb.setStyleSheet("font-size: 11px;")
            self.topic_checkboxes[topic] = cb
            topics_layout.addWidget(cb, row, col)
            col += 1
            if col > 1: 
                col = 0
                row += 1

        topics_scroll.setWidget(topics_widget)
        status_main_layout.addWidget(topics_scroll)
        self.status_group.setLayout(status_main_layout)
        
        header_layout.addWidget(self.status_group)

        self.summary_icon_btn = QPushButton("📑")
        self.summary_icon_btn.setFixedSize(65, 65)
        self.summary_icon_btn.setToolTip("View/Edit Introduction Summary")
        self.summary_icon_btn.setStyleSheet("""
            QPushButton { font-size: 30px; border-radius: 32px; background-color: #f8f9fa; border: 2px solid #ced4da; }
            QPushButton:hover { background-color: #e2e6ea; }
        """)
        self.summary_icon_btn.clicked.connect(self.show_summary_popup)
        header_layout.addWidget(self.summary_icon_btn, alignment=Qt.AlignmentFlag.AlignTop)

        right_layout.addLayout(header_layout) 

        # --- TEXT BOXES & NEW DROPDOWNS ---
        right_layout.addWidget(QLabel(" 🏷️ Project Name:"))
        self.project_name_input = QLineEdit()
        self.project_name_input.setPlaceholderText("Type or paste Project Name here...")
        self.project_name_input.setStyleSheet("background-color: white; font-size: 14px; padding: 5px; border: 1px solid #ced4da; border-radius: 4px;")
        right_layout.addWidget(self.project_name_input) 
        
        # Inject the new Popup Dropdowns
        dropdowns_data = [
            ("Scope Of Task Directive", ["Option 1", "Option 2", "Option 3"]),
            ("Stakeholders", ["Internal", "External", "Both"]),
            ("Certification Work Breakdown", ["Type A", "Type B", "Type C"]),
            ("Task Allocation", ["Auto", "Manual", "Hybrid"]),
            ("Communication Type", ["Email", "Meeting", "Report"])
        ]
        
        for label_text, items in dropdowns_data:
            row_layout = QHBoxLayout()
            lbl = QLabel(f" ⏷ {label_text}:")
            lbl.setStyleSheet("font-weight: bold; color: #0056b3;")
            lbl.setFixedWidth(200)
            combo = PopupMultiSelect(items)
            self.combos[label_text] = combo
            row_layout.addWidget(lbl)
            row_layout.addWidget(combo)
            right_layout.addLayout(row_layout)

        right_layout.addWidget(QLabel(" ✍️ Manual Notes, Stakeholders, & Other Data:"))
        self.manual_input = NotesEditor()
        self.manual_input.setPlaceholderText("Extract Stakeholders, Scope, and flowcharts here...")
        self.manual_input.setStyleSheet("background-color: white; font-size: 13px;")
        right_layout.addWidget(self.manual_input) 

        btn_layout = QHBoxLayout()
        self.save_docx_btn = QPushButton(" 📝 Generate Task Directive (Word)")
        self.save_docx_btn.setStyleSheet("background-color: #2b579a; color: white; font-weight: bold;")
        self.save_docx_btn.setFixedHeight(50)
        self.save_docx_btn.clicked.connect(self.export_docx)
        btn_layout.addWidget(self.save_docx_btn)
        
        right_layout.addLayout(btn_layout)

        self.splitter.addWidget(left_container)
        self.splitter.addWidget(self.right_container)
        self.splitter.setStretchFactor(0, 5) 
        self.splitter.setStretchFactor(1, 5) 
        layout.addWidget(self.splitter)

    # ------------------------------------------
    # --- EDITABLE SUMMARY POPUP ---
    # ------------------------------------------
    def show_summary_popup(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("📑 Edit Introduction Summary")
        dialog.resize(600, 400)
        layout = QVBoxLayout(dialog)

        text_edit = QTextEdit()
        text_edit.setText(self.intro_text_data)
        text_edit.setStyleSheet("font-size: 14px; line-height: 1.6; padding: 10px;")
        
        text_edit.textChanged.connect(lambda: setattr(self, 'intro_text_data', text_edit.toPlainText()))

        close_btn = QPushButton("Save & Close")
        close_btn.setFixedHeight(40)
        close_btn.clicked.connect(dialog.accept)

        layout.addWidget(text_edit)
        layout.addWidget(close_btn)
        dialog.exec()

    # ------------------------------------------
    # --- AUTO-PASTING FUNCTIONS ---
    # ------------------------------------------
    def add_extracted_text(self, text):
        QApplication.clipboard().setText(text) 
        self.manual_input.append(text + "\n")
        
    def add_extracted_image(self, pixmap):
        if pixmap.width() > 500:
            pixmap = pixmap.scaledToWidth(500, Qt.TransformationMode.SmoothTransformation)
        QApplication.clipboard().setPixmap(pixmap)
        cursor = self.manual_input.textCursor()
        cursor.movePosition(QTextCursor.MoveOperation.End) 
        self.manual_input.setTextCursor(cursor)
        cursor.insertImage(pixmap.toImage()) 
        self.manual_input.append("\n") 
        
    def generate_summary_from_selection(self, text):
        self.status_label.setText("Summarizing selection...")
        QApplication.processEvents()
        summary = simple_summarize(text, num_sentences=4)
        if summary:
            self.intro_text_data = summary 
            self.summary_icon_btn.setStyleSheet("""
                QPushButton { font-size: 30px; border-radius: 32px; background-color: #d4edda; border: 2px solid #28a745; }
                QPushButton:hover { background-color: #c3e6cb; }
            """)
            self.status_label.setText("Ready")
            self.topic_checkboxes["Introduction"].setChecked(True)
            QMessageBox.information(self, "Success", "Selection summarized! Click the 📑 icon to view or edit it.")
        else:
            self.status_label.setText("Ready")
            QMessageBox.warning(self, "Empty", "Could not generate a summary. Try selecting a larger paragraph.")

    # ------------------------------------------
    # --- UI INTERACTION FUNCTIONS ---
    # ------------------------------------------
    def eventFilter(self, source, event):
        if source == self.scroll_area.viewport() and event.type() == QEvent.Type.Wheel:
            if QApplication.keyboardModifiers() == Qt.KeyboardModifier.ControlModifier:
                delta = event.angleDelta().y()
                if delta > 0: self.zoom_in(step=50)
                elif delta < 0: self.zoom_out(step=50)
                return True 
        return super().eventFilter(source, event)

    def toggle_maximize(self):
        if not self.is_maximized:
            self.right_container.hide() 
            self.maximize_btn.setText("🗗") 
            self.is_maximized = True
        else:
            self.right_container.show() 
            self.maximize_btn.setText("🗖") 
            self.is_maximized = False

    def zoom_in(self, step=150):
        self.current_zoom += step
        self.refresh_pdf_view()

    def zoom_out(self, step=150):
        if self.current_zoom - step > 200: 
            self.current_zoom -= step
            self.refresh_pdf_view()

    def refresh_pdf_view(self):
        if not self.page_data: return

        if self.page_layout.count() != len(self.page_data):
            while self.page_layout.count():
                child = self.page_layout.takeAt(0)
                if child.widget(): child.widget().deleteLater()
            for pixmap, page_idx in self.page_data:
                page = self.doc.load_page(page_idx)
                lbl = PdfPageLabel(page, self.scroll_area)
                lbl.text_extracted.connect(self.add_extracted_text)
                lbl.image_extracted.connect(self.add_extracted_image)
                lbl.summary_requested.connect(self.generate_summary_from_selection)
                lbl.setStyleSheet("background-color: white; border: 1px solid black; margin-bottom: 10px;")
                self.page_layout.addWidget(lbl)

        for i in range(self.page_layout.count()):
            lbl = self.page_layout.itemAt(i).widget()
            if isinstance(lbl, PdfPageLabel) and i < len(self.page_data):
                scaled_pixmap = self.page_data[i][0].scaledToWidth(self.current_zoom, Qt.TransformationMode.SmoothTransformation)
                lbl.setPixmap(scaled_pixmap)
                lbl.setFixedSize(scaled_pixmap.size())

    # ------------------------------------------
    # --- CORE LOGIC & FUNCTIONS ---
    # ------------------------------------------
    def upload_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Document", "", "Documents (*.pdf *.txt)")
        if path:
            self.file_path = path
            
            self.check_upload.setChecked(False)
            for cb in self.topic_checkboxes.values():
                cb.setChecked(False)
                
            self.project_name_input.clear()
            self.intro_text_data = "" 
            self.summary_icon_btn.setStyleSheet("""
                QPushButton { font-size: 30px; border-radius: 32px; background-color: #f8f9fa; border: 2px solid #ced4da; }
                QPushButton:hover { background-color: #e2e6ea; }
            """)
            
            # Reset Combos
            for combo in self.combos.values():
                if combo.popup:
                    for i in range(combo.popup.list_widget.count()):
                        combo.popup.list_widget.item(i).setCheckState(Qt.CheckState.Unchecked)
                combo.setText("Select options...")
            
            self.manual_input.clear() 
            self.page_data.clear() 
            
            try:
                if path.lower().endswith('.pdf'):
                    self.text_preview.hide()
                    self.scroll_area.show()
                    self.doc = fitz.open(path)
                    
                    self.status_label.setText("Loading PDF Pages...")
                    QApplication.processEvents()

                    for i, page in enumerate(self.doc):
                        if i < 15: 
                            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
                            img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888).copy()
                            self.page_data.append((QPixmap.fromImage(img), i))
                    
                    self.current_zoom = 600 
                    self.refresh_pdf_view() 
                else:
                    self.scroll_area.hide()
                    self.text_preview.show()
                    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                        self.text_preview.setText(f.read())

                self.check_upload.setChecked(True)
                
                full_text = self.text_preview.toPlainText() if not self.doc else "".join(page.get_text() for page in self.doc)
                self.start_analysis_thread(full_text)
                
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Could not load file: {str(e)}")

    def start_analysis_thread(self, text):
        self.status_label.setText("Scanning document with Python NLP...")
        self.progress_bar.show()
        self.worker = AnalysisWorker(text, self.topics)
        self.worker.finished.connect(self.handle_analysis_done)
        self.worker.error.connect(lambda e: self.status_label.setText(f"Analysis Error: {e}"))
        self.worker.start()

    def handle_analysis_done(self, results):
        self.progress_bar.hide()
        self.status_label.setText("Analysis Complete")
        
        if results["project_name"]:
            self.project_name_input.setText(results["project_name"])
            
        if results["introduction"]:
            self.intro_text_data = results["introduction"]
            self.summary_icon_btn.setStyleSheet("""
                QPushButton { font-size: 30px; border-radius: 32px; background-color: #d4edda; border: 2px solid #28a745; }
                QPushButton:hover { background-color: #c3e6cb; }
            """)
            
        for topic in results["found_topics"]:
            if topic in self.topic_checkboxes:
                self.topic_checkboxes[topic].setChecked(True)

    # ------------------------------------------
    # --- TEMPLATE INJECTION & EXPORT ---
    # ------------------------------------------
    def export_docx(self):
        save_path, _ = QFileDialog.getSaveFileName(self, "Save Word", "Task_Directive_Final.docx", "Word (*.docx)")
        if not save_path: return
            
        try:
            def add_page_border(section):
                sectPr = section._sectPr
                pgBorders = OxmlElement('w:pgBorders')
                pgBorders.set(qn('w:offsetFrom'), 'page')
                for side in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement(f'w:{side}')
                    border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '12'); border.set(qn('w:space'), '24'); border.set(qn('w:color'), '000000')
                    pgBorders.append(border)
                sectPr.append(pgBorders)

            def remove_page_border(section):
                sectPr = section._sectPr
                for el in sectPr.findall(qn('w:pgBorders')): sectPr.remove(el)

            def add_heading(doc, text):
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True; run.font.size = Pt(12)

            doc = Document()
            style = doc.styles['Normal']
            style.font.name = 'Calibri'; style.font.size = Pt(11)

            # --- SECTION 1: COVER PAGE ---
            section1 = doc.sections[0]
            section1.top_margin = Inches(1.5); section1.bottom_margin = Inches(1.5)
            section1.left_margin = Inches(1.5); section1.right_margin = Inches(1.5)
            add_page_border(section1)

            doc.add_paragraph("File No. ________________________")
            p = doc.add_paragraph()
            run = p.add_run("\nTASK DIRECTIVE\n"); run.bold = True; run.font.size = Pt(20)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph("________ /2026"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Issue No."; table.cell(0, 1).text = "Date of Issue:"

            p_name = self.project_name_input.text().strip()
            if not p_name or p_name.lower() == "not found":
                p_name = "__________________________________________"
            doc.add_paragraph(f"\nPROJECT NAME: {p_name}")

            p = doc.add_paragraph("\n\n[ LOGO HERE ]"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("ABC Organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Address of organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Organization details").alignment = WD_ALIGN_PARAGRAPH.CENTER

            # --- SECTION 2: MAIN CONTENT ---
            section2 = doc.add_section(WD_SECTION.NEW_PAGE)
            section2.header.is_linked_to_previous = False; section2.footer.is_linked_to_previous = False
            section2.top_margin = Inches(0.75); section2.bottom_margin = Inches(0.75)
            section2.left_margin = Inches(0.75); section2.right_margin = Inches(0.75)
            remove_page_border(section2)

            add_heading(doc, "1. Introduction [Automated]")
            intro_text = self.intro_text_data.strip()
            if not intro_text: intro_text = "__________________________________________________"
            doc.add_paragraph(intro_text)

            add_heading(doc, "2. Reference [Default]")
            doc.add_paragraph("__________________________________________________")
            add_heading(doc, "3. Basis Of Task Directive [Default]")
            doc.add_paragraph("__________________________________________________")
            
            # Extract from Combo boxes
            scope_combo = self.combos.get("Scope Of Task Directive")
            scope_text = "\n".join(scope_combo.checkedItems()) if scope_combo and scope_combo.checkedItems() else "____"
            
            add_heading(doc, "4. Scope Of Task Directive [Drop down menu]")
            doc.add_paragraph(f"To assign the certification respectively:\n{scope_text}\n\n1) ____\n2) ____\n3) ____")

            sth_combo = self.combos.get("Stakeholders")
            sth_text = "\n".join(sth_combo.checkedItems()) if sth_combo and sth_combo.checkedItems() else "Both"
            
            add_heading(doc, f"5. Stakeholders [Automated + Manual Notes] ({sth_text})")
            doc.add_paragraph("The following are the major stakeholders and manual notes extracted:")
            doc.add_paragraph(self.manual_input.toPlainText()) 

            table = doc.add_table(rows=4, cols=4); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No.", "Organisation", "Role", "Activities"]): table.cell(0, i).text = h

            cert_combo = self.combos.get("Certification Work Breakdown")
            cert_text = "\n".join(cert_combo.checkedItems()) if cert_combo and cert_combo.checkedItems() else "______________________________________________"
            add_heading(doc, "6. Certification Work Breakdown [Drop down menu]")
            doc.add_paragraph(cert_text)
            
            task_combo = self.combos.get("Task Allocation")
            task_text = "\n".join(task_combo.checkedItems()) if task_combo and task_combo.checkedItems() else "______________________________________________"
            add_heading(doc, "7. Task Allocation [Default]")
            doc.add_paragraph(task_text)
            add_heading(doc, "7.1 Coordinating Directorate [Default]")
            doc.add_paragraph("______________________________________________")
            add_heading(doc, "7.2 Single Point of Contact (SPoC) [Default]")
            doc.add_paragraph("______________________________________________")

            add_heading(doc, "7.3 Certification Task Allocation")
            table = doc.add_table(rows=4, cols=4); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No.", "Certification Activity", "Certification Work Centre", "Responsible Head"]): table.cell(0, i).text = h

            add_heading(doc, "7.4 Issue of Clearance [Default]")
            doc.add_paragraph("a. ___\nb. ___\nc. ___\nd. ___\ne. ___\nf. ___\ng. ___")

            add_heading(doc, "8. SCRB And TARB [Default]")
            doc.add_paragraph("______________________________________________")
            
            comm_combo = self.combos.get("Communication Type")
            comm_text = "\n".join(comm_combo.checkedItems()) if comm_combo and comm_combo.checkedItems() else "______________________________________________"
            add_heading(doc, "9. Communication [Default]")
            doc.add_paragraph(comm_text)
            add_heading(doc, "10. Certification Progress Review [Default]")
            doc.add_paragraph("______________________________________________\n\n(__________)")

            add_heading(doc, "11. Distribution List")
            doc.add_paragraph("11.1 External Organization\n1. ___\n2. ___\n3. ___\n11.2 Internal Distribution")

            # --- ANNEXURES ---
            doc.add_page_break()
            add_heading(doc, "Annexure-1")
            doc.add_paragraph("Work Assignment List of LRUs [Automate + Manual]")
            table = doc.add_table(rows=4, cols=6); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No", "ABC", "Abc2", "Abc3", "Abc4", "Abc5"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Integration Checks and Clearance [Manual]")
            table = doc.add_table(rows=3, cols=5); table.style = 'Table Grid'
            for i, h in enumerate(["Sl no.", "Abc1", "Abc2", "Abc3", "Abc4"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Annexure-2")
            doc.add_paragraph("Contact details of dealing officers and RDs [Manual]")
            table = doc.add_table(rows=3, cols=5); table.style = 'Table Grid'
            for i, h in enumerate(["Sl no.", "Abc1", "Abc2", "Abc3", "Abc4"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Annexure-3")
            doc.add_paragraph("Product Break Down Structure [Automate]")

            doc.save(save_path)
            QMessageBox.information(self, "Success", "Task Directive Template created successfully!")

        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to build template: {str(e)}")

if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = OfflineApp()
    window.show()
    sys.exit(app.exec())
