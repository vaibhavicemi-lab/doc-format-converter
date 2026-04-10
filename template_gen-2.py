import sys
import os
import re
import fitz  # PyMuPDF
import pandas as pd
from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout,
    QHBoxLayout, QPushButton, QTextEdit, QFileDialog,
    QLabel, QSplitter, QMessageBox, QScrollArea, QProgressBar,
    QCheckBox, QGroupBox, QDialog, QRubberBand, QMenu, QLineEdit,
    QGridLayout, QFrame, QComboBox, QInputDialog, QStyleFactory
)
from PyQt6.QtGui import QPixmap, QImage, QFont, QTextDocument, QTextCursor, QPageLayout, QStandardItemModel, QStandardItem,QPalette, QColor



# ==========================================
# --- CHECKABLE COMBO BOX ---
# ==========================================
class CheckableComboBox(QComboBox):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.view().pressed.connect(self.handleItemPressed)
        self._model = QStandardItemModel(self)
        self.setModel(self._model)
        self._ignore_hide = False
        self.setModelColumn(0)
        self.setStyleSheet(
            "QComboBox QAbstractItemView { background-color: #ffffff; color: #000000; }"
        )

    def handleItemPressed(self, index):
        item = self._model.itemFromIndex(index)
        if item is None:
            return
        if item.checkState() == Qt.CheckState.Checked:
            item.setCheckState(Qt.CheckState.Unchecked)
        else:
            item.setCheckState(Qt.CheckState.Checked)
        self._ignore_hide = True
        self._model.dataChanged.emit(index, index)

    def addItem(self, text, checked=False):
        item = QStandardItem(text)
        item.setFlags(
            Qt.ItemFlag.ItemIsSelectable |
            Qt.ItemFlag.ItemIsEnabled |
            Qt.ItemFlag.ItemIsUserCheckable
        )
        item.setCheckState(Qt.CheckState.Checked if checked else Qt.CheckState.Unchecked)
        self._model.appendRow(item)
        
    def addItems(self, texts):
        for text in texts:
            self.addItem(text)

    def checkedItems(self):
        checked_items = []
        for i in range(self._model.rowCount()):
            item = self._model.item(i)
            if item.checkState() == Qt.CheckState.Checked and item.text() != "Enter Other...":
                checked_items.append(item.text())
        return checked_items

    def hidePopup(self):
        if self._ignore_hide:
            self._ignore_hide = False
            return
        super().hidePopup()

    def currentText(self):
        checked = self.checkedItems()
        return ", ".join(checked) if checked else ""
from PyQt6.QtCore import Qt, QThread, pyqtSignal, QEvent, QRect, QPoint, QSize, QMarginsF
from PyQt6.QtPrintSupport import QPrinter

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==========================================
# --- PURE PYTHON SUMMARIZER ---
# ==========================================
def simple_summarize(text, num_sentences=4):
    if not text: return ""
    stop_words = {"the", "is", "in", "and", "to", "of", "a", "for", "on", "with", "as", "by", "this", "that", "it", "are", "be", "or", "an", "at", "from", "which", "will"}
    sentences = re.split(r'(?<=[.!?]) +|\n+', text)
    sentences = [s.strip() for s in sentences if len(s.strip()) > 5]
    if len(sentences) <= num_sentences: return " ".join(sentences)
        
    words = re.findall(r'\b\w+\b', text.lower())
    freq = {}
    for w in words:
        if w not in stop_words and not w.isnumeric(): freq[w] = freq.get(w, 0) + 1
            
    scores = {}
    for i, s in enumerate(sentences):
        score = 0
        s_words = re.findall(r'\b\w+\b', s.lower())
        for w in s_words:
            if w in freq: score += freq[w]
        scores[i] = score / max(len(s_words), 1)
        
    top_indices = sorted(sorted(scores, key=scores.get, reverse=True)[:num_sentences])
    return " ".join([sentences[i] for i in top_indices])

# ==========================================
# --- SMART TEXT EDITOR (SUPPORTS IMAGES)---
# ==========================================
class NotesEditor(QTextEdit):
    def canInsertFromMimeData(self, source):
        if source.hasImage() or source.hasText(): return True
        return super().canInsertFromMimeData(source)

    def insertFromMimeData(self, source):
        if source.hasImage():
            clipboard = QApplication.clipboard()
            image = clipboard.image() 
            if not image.isNull():
                cursor = self.textCursor()
                if image.width() > 500:
                    image = image.scaledToWidth(500, Qt.TransformationMode.SmoothTransformation)
                cursor.insertImage(image)
                self.insertPlainText("\n")
                return 
        if source.hasText():
            self.insertPlainText(source.text())
            return
        super().insertFromMimeData(source)

# ==========================================
# --- PURE CODE ANALYSIS WORKER ---
# ==========================================
class AnalysisWorker(QThread):
    finished = pyqtSignal(dict) 
    error = pyqtSignal(str)

    def __init__(self, text, topics):
        super().__init__()
        self.text = text
        self.topics = topics

    def run(self):
        try:
            results = {
                "project_name": "",
                "introduction": "",
                "found_topics": []
            }

            match = re.search(r'Project\s*Name[\s:]*(.+)', self.text, re.IGNORECASE)
            if match:
                clean_name = match.group(1).strip()
                if not clean_name.startswith("___"):
                    results["project_name"] = clean_name[:60] 

            results["introduction"] = simple_summarize(self.text, num_sentences=4)

            for topic in self.topics:
                if topic.lower() in self.text.lower():
                    results["found_topics"].append(topic)

            if results["introduction"] and "Introduction" not in results["found_topics"]:
                results["found_topics"].append("Introduction")

            self.finished.emit(results)
        except Exception as e:
            self.error.emit(str(e))

# ==========================================
# --- INTERACTIVE PDF SNIPPING TOOL ---
# ==========================================
class PdfPageLabel(QLabel):
    text_extracted = pyqtSignal(str)
    image_extracted = pyqtSignal(QPixmap)
    summary_requested = pyqtSignal(str) 

    def __init__(self, page, scroll_area, parent=None):
        super().__init__(parent)
        self.page = page  
        self.scroll_area = scroll_area 
        self.rubber_band = QRubberBand(QRubberBand.Shape.Rectangle, self)
        self.origin = QPoint()
        self.pan_start_pos = None

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            self.origin = event.pos()
            self.rubber_band.setGeometry(QRect(self.origin, QSize()))
            self.rubber_band.show()
        elif event.button() in [Qt.MouseButton.RightButton, Qt.MouseButton.MiddleButton]:
            self.pan_start_pos = event.globalPosition().toPoint()
            self.setCursor(Qt.CursorShape.ClosedHandCursor)

    def mouseMoveEvent(self, event):
        if not self.origin.isNull() and (event.buttons() & Qt.MouseButton.LeftButton):
            self.rubber_band.setGeometry(QRect(self.origin, event.pos()).normalized())
            
        elif self.pan_start_pos is not None and (event.buttons() & (Qt.MouseButton.RightButton | Qt.MouseButton.MiddleButton)):
            current_pos = event.globalPosition().toPoint()
            delta = current_pos - self.pan_start_pos
            
            h_bar = self.scroll_area.horizontalScrollBar()
            v_bar = self.scroll_area.verticalScrollBar()
            
            h_bar.setValue(h_bar.value() - delta.x())
            v_bar.setValue(v_bar.value() - delta.y())
            
            self.pan_start_pos = current_pos

    def mouseReleaseEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton:
            rect = self.rubber_band.geometry()
            if rect.width() > 10 and rect.height() > 10:
                self.show_context_menu(event.pos(), rect)
            else:
                self.rubber_band.hide()
            self.origin = QPoint()
            
        elif event.button() in [Qt.MouseButton.RightButton, Qt.MouseButton.MiddleButton]:
            self.unsetCursor()
            self.pan_start_pos = None
            
    def get_text_from_rect(self, rect):
        if self.pixmap() and self.page:
            scale_x = self.page.rect.width / self.pixmap().width()
            scale_y = self.page.rect.height / self.pixmap().height()
            pdf_rect = fitz.Rect(rect.left() * scale_x, rect.top() * scale_y, rect.right() * scale_x, rect.bottom() * scale_y)
            pdf_rect = pdf_rect + (-3, -3, 3, 3)
            return self.page.get_textbox(pdf_rect).strip()
        return ""

    def show_context_menu(self, pos, rect):
        menu = QMenu(self)
        copy_img_action = menu.addAction("Extract as Image")
        copy_text_action = menu.addAction("Extract Text to Notes")
        menu.addSeparator()
        summarize_action = menu.addAction("Summarize Selection (Set as Intro)")

        action = menu.exec(self.mapToGlobal(pos))
        
        if action == copy_img_action:
            if self.pixmap(): self.image_extracted.emit(self.pixmap().copy(rect)) 
        elif action == copy_text_action:
            text = self.get_text_from_rect(rect)
            if text: self.text_extracted.emit(text) 
        elif action == summarize_action:
            text = self.get_text_from_rect(rect)
            if text: self.summary_requested.emit(text)
            
        self.rubber_band.hide()

# ==========================================
# --- MULTI SELECT COMPONENT ---
# ==========================================
class MultiSelectWidget(QWidget):
    def __init__(self, options, parent=None):
        super().__init__(parent)
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(5)
        
        self.display_box = QTextEdit()
        self.display_box.setReadOnly(True)
        self.display_box.setFixedHeight(65)
        self.display_box.setPlaceholderText("Selected options will appear here...")
        self.display_box.setStyleSheet("background-color: #f1f5f9; color: #0f172a; border: 1px solid #cbd5e1; border-radius: 6px;")
        
        self.combo = CheckableComboBox()
        clean_opts = [o for o in options if o != "Enter Other..."]
        self.combo.addItems(clean_opts)
        
        input_layout = QHBoxLayout()
        self.custom_input = QLineEdit()
        self.custom_input.setPlaceholderText("Or type other option here...")
        self.add_btn = QPushButton("Add")
        self.add_btn.setObjectName("ToolBtn")
        self.add_btn.setFixedWidth(60)
        self.add_btn.clicked.connect(self.add_custom_option)
        self.custom_input.returnPressed.connect(self.add_custom_option)
        
        input_layout.addWidget(self.custom_input)
        input_layout.addWidget(self.add_btn)
        
        layout.addWidget(self.combo)
        layout.addWidget(self.display_box)
        layout.addLayout(input_layout)
        
        self.combo._model.dataChanged.connect(self.update_display)
        
    def add_custom_option(self):
        text = self.custom_input.text().strip()
        if text:
            existing = [self.combo._model.item(i).text() for i in range(self.combo._model.rowCount())]
            if text not in existing:
                self.combo.addItem(text, checked=True)
            else:
                for i in range(self.combo._model.rowCount()):
                    if self.combo._model.item(i).text() == text:
                        self.combo._model.item(i).setCheckState(Qt.CheckState.Checked)
            self.custom_input.clear()
            self.update_display()
            
    def update_display(self):
        self.display_box.setText("\\n".join(self.combo.checkedItems()))
        
    def checkedItems(self):
        return self.combo.checkedItems()


# ==========================================
# --- MAIN APPLICATION ---
# ==========================================
class OfflineApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Offline Document Assistant - Fast Analyzer")
        self.resize(1450, 950) 
        
        self.file_path = ""
        self.doc = None
        self.page_data = [] 
        self.current_zoom = 600 
        self.is_maximized = False
        self.original_file_path = ""
        
        self.intro_text_data = "" 
        
        self.apply_stylesheet()
        self.init_ui()

    def apply_stylesheet(self):
        self.setStyleSheet("""
            QWidget { font-family: 'Segoe UI', sans-serif; font-size: 11pt; }
            QMainWindow, #mainWidget, QDialog, QMessageBox { background-color: #f8fafc; color: #0f172a; }
            #leftContainer { background: #ffffff; border-right: 1px solid #e2e8f0; }
            #rightContainer { background: #f1f5f9; }
            QFrame[class="Card"] { background: white; border-radius: 12px; border: 1px solid #e2e8f0; color: #0f172a; }
            QPushButton { background: #4f46e5; color: white; border: none; padding: 12px; border-radius: 8px; font-weight: bold; }
            QPushButton:hover { background: #4338ca; }
            QPushButton#ActionBtn { height: 40px; font-weight: bold; background-color: #f1f5f9; color: #334155; padding: 0 15px; border: 1px solid #cbd5e1; }
            QPushButton#ActionBtn:hover { background-color: #e2e8f0; }
            QPushButton#ToolBtn { background: white; color: #334155; border: 1px solid #cbd5e1; font-size: 16px; font-weight: bold; border-radius: 4px; padding: 5px; }
            QPushButton#ToolBtn:hover { background: #f1f5f9; }
            QPushButton#IconBtn { font-size: 24px; border-radius: 25px; background-color: #f8f9fa; border: 2px solid #ced4da; color: black; padding: 0; }
            QPushButton#IconBtn:hover { background-color: #e2e6ea; }
            QPushButton#IconBtnActive { font-size: 24px; border-radius: 25px; background-color: #d4edda; border: 2px solid #28a745; color: black; padding: 0; }
            QGroupBox { background: transparent; border: none; color: #0f172a; }
            QComboBox { padding: 8px; border-radius: 6px; border: 1px solid #cbd5e1; background-color: #ffffff; color: #000000; }
            QComboBox::drop-down { border: 0px; }
            QComboBox QAbstractItemView { border: 1px solid #cbd5e1; background-color: #ffffff; color: #000000; selection-background-color: #e2e8f0; selection-color: #000000; outline: 0; }
            QComboBox QAbstractItemView::item { background-color: #ffffff; color: #000000; padding: 8px; }
            QComboBox QAbstractItemView::item:selected { background-color: #e2e8f0; color: #000000; }
            QSplitter::handle { background-color: #94a3b8; border-radius: 2px; margin: 2px; }
            QLabel { color: #0f172a; background: transparent; }
            QLabel[class="Heading"] { font-size: 18px; font-weight: bold; color: #4f46e5; margin: 0; background: transparent; }
            QLineEdit, QTextEdit { background-color: #ffffff; color: #000000; border: 1px solid #cbd5e1; border-radius: 6px; padding: 8px; }
            QCheckBox { color: #0f172a; background: transparent; }
            QProgressBar { border: 1px solid #cbd5e1; background: #e2e8f0; border-radius: 5px; text-align: center; color: #0f172a; }
            QProgressBar::chunk { background-color: #4f46e5; }
            QScrollArea { border: none; background: transparent; }
            QScrollArea > QWidget > QWidget { background: transparent; }
            QMenu { background-color: white; color: #0f172a; border: 1px solid #cbd5e1; }
            QMenu::item:selected { background-color: #f1f5f9; }
            QToolTip { background: white; color: #0f172a; border: 1px solid #cbd5e1; }
        """)

    def init_ui(self):
        main_widget = QWidget()
        main_widget.setObjectName("mainWidget")
        self.setCentralWidget(main_widget)
        layout = QVBoxLayout(main_widget)
        layout.setContentsMargins(0, 0, 0, 0)
        
        self.splitter = QSplitter(Qt.Orientation.Horizontal)
        self.splitter.setHandleWidth(6)

        # --- LEFT SIDE: Document Viewer ---
        left_container = QWidget()
        left_container.setObjectName("leftContainer")
        left_layout = QVBoxLayout(left_container)
        left_layout.setContentsMargins(20, 20, 20, 20)
        left_layout.setSpacing(15)
        
        left_card = QFrame()
        left_card.setProperty("class", "Card")
        left_card_layout = QVBoxLayout(left_card)
        left_card_layout.setContentsMargins(15, 15, 15, 15)
        left_card_layout.setSpacing(10)

        left_controls_layout = QHBoxLayout()
        
        self.upload_btn = QPushButton("Upload PDF Document")
        self.upload_btn.setObjectName("ActionBtn")
        self.upload_btn.clicked.connect(self.upload_file)
        
        self.close_preview_btn = QPushButton("Close Preview")
        self.close_preview_btn.setObjectName("ActionBtn")
        self.close_preview_btn.setStyleSheet("color: #dc2626;")
        self.close_preview_btn.clicked.connect(self.close_preview)
        self.close_preview_btn.hide()
        
        self.zoom_out_btn = QPushButton("-")
        self.zoom_out_btn.setObjectName("ToolBtn")
        self.zoom_in_btn = QPushButton("+")
        self.zoom_in_btn.setObjectName("ToolBtn")
        self.maximize_btn = QPushButton("🗖") 
        self.maximize_btn.setObjectName("ToolBtn")
        for btn in [self.zoom_out_btn, self.zoom_in_btn, self.maximize_btn]:
            btn.setFixedSize(40, 40)
            
        self.zoom_out_btn.clicked.connect(lambda: self.zoom_out(150))
        self.zoom_in_btn.clicked.connect(lambda: self.zoom_in(150))
        self.maximize_btn.clicked.connect(self.toggle_maximize)

        left_controls_layout.addWidget(self.upload_btn)
        left_controls_layout.addWidget(self.close_preview_btn)
        left_controls_layout.addStretch()
        z_label = QLabel("Zoom:")
        
        left_controls_layout.addWidget(z_label)
        left_controls_layout.addWidget(self.zoom_out_btn)
        left_controls_layout.addWidget(self.zoom_in_btn)
        left_controls_layout.addWidget(self.maximize_btn)

        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setStyleSheet("QScrollArea { border: none; background: #f1f3f8; border-radius: 10px; }") 
        self.page_container = QWidget()
        self.page_container.setStyleSheet("background: #f1f3f8;")
        self.page_layout = QVBoxLayout(self.page_container)
        self.page_layout.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignHCenter) 
        self.scroll_area.setWidget(self.page_container)
        self.scroll_area.viewport().installEventFilter(self)
        self.scroll_area.verticalScrollBar().setSingleStep(30)
        self.scroll_area.horizontalScrollBar().setSingleStep(30)

        self.text_preview = QTextEdit()
        self.text_preview.setReadOnly(True)
        self.text_preview.hide()
        self.text_preview.setStyleSheet("border: none; background: #f1f3f8; border-radius: 10px; padding: 10px; color: #333333;")

        left_card_layout.addLayout(left_controls_layout)
        tip_label = QLabel("Click and drag on the PDF to select and extract text, flowcharts, or generate your Introduction Summary!")
        tip_label.setStyleSheet("font-style: italic;")
        left_card_layout.addWidget(tip_label)
        left_card_layout.addWidget(self.scroll_area)
        left_card_layout.addWidget(self.text_preview)

        left_layout.addWidget(left_card)

        # --- RIGHT SIDE: Tools & Extraction ---
        self.right_container = QScrollArea()
        self.right_container.setWidgetResizable(True)
        self.right_container.setObjectName("rightContainer")
        self.right_container.setStyleSheet("QScrollArea#rightContainer { border: none; background: transparent; }")
        self.right_container.verticalScrollBar().setSingleStep(30)
        
        right_widget = QWidget()
        right_widget.setStyleSheet("background: transparent;")
        right_layout = QVBoxLayout(right_widget)
        right_layout.setContentsMargins(20, 20, 20, 20)
        right_layout.setSpacing(20)
        self.right_container.setWidget(right_widget)

        # Card 1: Document Status & Checklist
        status_card = QFrame()
        status_card.setProperty("class", "Card")
        status_main_layout = QVBoxLayout(status_card)
        status_main_layout.setContentsMargins(15, 15, 15, 15)
        
        top_bar_layout = QHBoxLayout()
        title_lbl = QLabel("<h3>Task Directive Extraction Checklist</h3>")
        title_lbl.setProperty("class", "Heading")
        top_bar_layout.addWidget(title_lbl)
        top_bar_layout.addStretch()
        
        # Theme toggle removed entirely
        
        self.summary_icon_btn = QPushButton("AI")
        self.summary_icon_btn.setFixedSize(50, 50)
        self.summary_icon_btn.setObjectName("IconBtn")
        self.summary_icon_btn.setToolTip("View/Edit Introduction Summary")
        self.summary_icon_btn.clicked.connect(self.show_summary_popup)
        top_bar_layout.addWidget(self.summary_icon_btn)
        status_main_layout.addLayout(top_bar_layout)
        
        chk_layout = QHBoxLayout()
        self.check_upload = QCheckBox("File Uploaded")
        self.check_upload.setEnabled(False)
        self.check_upload.setStyleSheet("font-weight: bold;")
        chk_layout.addWidget(self.check_upload)
        
        self.topics = [
            "Introduction", "Reference", "Basis Of Task Directive",
            "Scope Of Task Directive", "Stakeholders", "Certification Work Breakdown",
            "Task Allocation", "Coordinating Directorate", "Single Point of Contact",
            "Certification Task Allocation", "Issue of Clearance", "SCRB And TARB",
            "Communication", "Certification Progress Review", "Distribution List"
        ]
        self.topic_checkboxes = {}
        
        topics_scroll = QScrollArea()
        topics_scroll.setWidgetResizable(True)
        topics_scroll.setFixedHeight(120) 
        topics_scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        
        topics_widget = QWidget()
        topics_layout = QGridLayout(topics_widget) 
        topics_layout.setContentsMargins(5, 5, 5, 5)
        
        row, col = 0, 0
        for topic in self.topics:
            cb = QCheckBox(topic)
            
            self.topic_checkboxes[topic] = cb
            topics_layout.addWidget(cb, row, col)
            col += 1
            if col > 1: 
                col = 0
                row += 1

        topics_scroll.setWidget(topics_widget)
        status_main_layout.addLayout(chk_layout)
        status_main_layout.addWidget(topics_scroll)

        right_layout.addWidget(status_card)

        # Card 2: Manual Data
        manual_card = QFrame()
        manual_card.setProperty("class", "Card")
        manual_layout = QVBoxLayout(manual_card)
        manual_layout.setContentsMargins(15, 15, 15, 15)

        manual_title = QLabel("<h3>Manual Data</h3>")
        manual_title.setProperty("class", "Heading")
        manual_layout.addWidget(manual_title)

        proj_label = QLabel("Project Name:")
        proj_label.setStyleSheet("font-weight: bold;")
        manual_layout.addWidget(proj_label)
        
        self.project_name_input = QLineEdit()
        self.project_name_input.setPlaceholderText("Type or paste Project Name here...")
        
        manual_layout.addWidget(self.project_name_input)

        dropdowns = [
            ("Scope Of Task Directive", ["Option 1", "Option 2", "Option 3", "Enter Other..."]),
            ("Stakeholders", ["Internal", "External", "Both", "Enter Other..."]),
            ("Certification Work Breakdown", ["Type A", "Type B", "Type C", "Enter Other..."]),
            ("Task Allocation", ["Auto", "Manual", "Hybrid", "Enter Other..."]),
            ("Communication Type", ["Email", "Meeting", "Report", "Enter Other..."])
        ]
        
        self.combos = []
        for label, items in dropdowns:
            row = QHBoxLayout()
            lbl = QLabel(label)
            lbl.setFixedWidth(200)
            lbl.setAlignment(Qt.AlignmentFlag.AlignTop | Qt.AlignmentFlag.AlignLeft)
            lbl.setStyleSheet("margin-top: 5px;")
            combo = MultiSelectWidget(items)
            row.addWidget(lbl)
            row.addWidget(combo)
            manual_layout.addLayout(row)
            self.combos.append(combo)
            
        self.scope_dropdown, self.stakeholder_dropdown, self.cert_dropdown, self.task_dropdown, self.comm_dropdown = self.combos

        notes_label = QLabel("Manual Notes, Stakeholders, & Other Data:")
        notes_label.setStyleSheet("font-weight: bold; margin-top: 10px;")
        manual_layout.addWidget(notes_label)
        
        self.manual_input = NotesEditor()
        self.manual_input.setPlaceholderText("Extract Stakeholders, Scope, and flowcharts here...")
        
        manual_layout.addWidget(self.manual_input)

        right_layout.addWidget(manual_card, stretch=1)

        # Card 3: Actions
        actions_card = QFrame()
        actions_card.setProperty("class", "Card")
        actions_layout = QVBoxLayout(actions_card)
        actions_layout.setContentsMargins(15, 15, 15, 15)
        
        actions_title = QLabel("<h3>Actions</h3>")
        actions_title.setProperty("class", "Heading")
        actions_layout.addWidget(actions_title)

        self.status_label = QLabel("Ready")
        self.status_label.setStyleSheet("font-style: italic;")
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0) 
        self.progress_bar.hide()
        
        actions_layout.addWidget(self.status_label)
        actions_layout.addWidget(self.progress_bar)

        btn_layout = QHBoxLayout()
        btn_layout.setSpacing(10)
        
        self.save_docx_btn = QPushButton("Generate Task Directive (Word)")
        self.save_docx_btn.setFixedHeight(45)
        self.save_docx_btn.setStyleSheet("background-color: #059669; color: white;")
        self.save_docx_btn.clicked.connect(self.export_docx)
        
        self.preview_pdf_btn = QPushButton("Preview PDF")
        self.preview_pdf_btn.setFixedHeight(45)
        self.preview_pdf_btn.setStyleSheet("background-color: #059669; color: white;")
        self.preview_pdf_btn.clicked.connect(self.preview_pdf)

        btn_layout.addWidget(self.save_docx_btn)
        btn_layout.addWidget(self.preview_pdf_btn)
        actions_layout.addLayout(btn_layout)

        right_layout.addWidget(actions_card)

        self.splitter.addWidget(left_container)
        self.splitter.addWidget(self.right_container)
        self.splitter.setStretchFactor(0, 5) 
        self.splitter.setStretchFactor(1, 5) 
        layout.addWidget(self.splitter)

    # ------------------------------------------
    # --- EDITABLE SUMMARY POPUP ---
    # ------------------------------------------
    def show_summary_popup(self):
        """Opens a popup window to view and edit the generated summary."""
        dialog = QDialog(self)
        dialog.setWindowTitle("Edit Introduction Summary")
        dialog.resize(600, 400)
        layout = QVBoxLayout(dialog)

        text_edit = QTextEdit()
        text_edit.setStyleSheet("background-color: white; color: black; font-size: 11pt;")
        text_edit.setText(self.intro_text_data)
        
        text_edit.textChanged.connect(lambda: setattr(self, 'intro_text_data', text_edit.toPlainText()))

        close_btn = QPushButton("Save & Close")
        close_btn.setFixedHeight(40)
        close_btn.clicked.connect(dialog.accept)

        layout.addWidget(text_edit)
        layout.addWidget(close_btn)
        dialog.exec()

    # ------------------------------------------
    # --- AUTO-PASTING FUNCTIONS ---
    # ------------------------------------------
    def add_extracted_text(self, text):
        QApplication.clipboard().setText(text) 
        self.manual_input.append(text + "\\n")
        
    def add_extracted_image(self, pixmap):
        if pixmap.width() > 500:
            pixmap = pixmap.scaledToWidth(500, Qt.TransformationMode.SmoothTransformation)
        QApplication.clipboard().setPixmap(pixmap)
        cursor = self.manual_input.textCursor()
        cursor.movePosition(QTextCursor.MoveOperation.End) 
        self.manual_input.setTextCursor(cursor)
        cursor.insertImage(pixmap.toImage()) 
        self.manual_input.append("\\n") 
        
    def generate_summary_from_selection(self, text):
        self.status_label.setText("Summarizing selection...")
        QApplication.processEvents()
        summary = simple_summarize(text, num_sentences=4)
        if summary:
            self.intro_text_data = summary 
            
            self.summary_icon_btn.setObjectName("IconBtnActive"); self.summary_icon_btn.style().unpolish(self.summary_icon_btn); self.summary_icon_btn.style().polish(self.summary_icon_btn)
            
            self.status_label.setText("Ready")
            self.topic_checkboxes["Introduction"].setChecked(True)
            QMessageBox.information(self, "Success", "Selection summarized! Click the AI icon to view or edit it.")
        else:
            self.status_label.setText("Ready")
            QMessageBox.warning(self, "Empty", "Could not generate a summary. Try selecting a larger paragraph.")

    # ------------------------------------------
    # --- UI INTERACTION FUNCTIONS ---
    # ------------------------------------------
    def eventFilter(self, source, event):
        if source == self.scroll_area.viewport() and event.type() == QEvent.Type.Wheel:
            if QApplication.keyboardModifiers() == Qt.KeyboardModifier.ControlModifier:
                delta = event.angleDelta().y()
                if delta > 0: self.zoom_in(step=50)
                elif delta < 0: self.zoom_out(step=50)
                return True 
        return super().eventFilter(source, event)

    def toggle_maximize(self):
        if not self.is_maximized:
            self.right_container.hide() 
            self.maximize_btn.setText("🗗") 
            self.is_maximized = True
        else:
            self.right_container.show() 
            self.maximize_btn.setText("🗖") 
            self.is_maximized = False

    def zoom_in(self, step=150):
        self.current_zoom += step
        self.refresh_pdf_view()

    def zoom_out(self, step=150):
        if self.current_zoom - step > 200: 
            self.current_zoom -= step
            self.refresh_pdf_view()

    def refresh_pdf_view(self):
        if not self.page_data: return

        if self.page_layout.count() != len(self.page_data):
            while self.page_layout.count():
                child = self.page_layout.takeAt(0)
                if child.widget(): child.widget().deleteLater()
            for pixmap, page_idx in self.page_data:
                page = self.doc.load_page(page_idx)
                lbl = PdfPageLabel(page, self.scroll_area)
                lbl.text_extracted.connect(self.add_extracted_text)
                lbl.image_extracted.connect(self.add_extracted_image)
                lbl.summary_requested.connect(self.generate_summary_from_selection)
                lbl.setStyleSheet("background-color: white; border: 1px solid #e0e0e0; margin-bottom: 10px; border-radius: 4px;")
                self.page_layout.addWidget(lbl)

        for i in range(self.page_layout.count()):
            lbl = self.page_layout.itemAt(i).widget()
            if isinstance(lbl, PdfPageLabel) and i < len(self.page_data):
                scaled_pixmap = self.page_data[i][0].scaledToWidth(self.current_zoom, Qt.TransformationMode.SmoothTransformation)
                lbl.setPixmap(scaled_pixmap)
                lbl.setFixedSize(scaled_pixmap.size())

    # ------------------------------------------
    # --- CORE LOGIC & FUNCTIONS ---
    # ------------------------------------------
    def close_preview(self):
        if not hasattr(self, 'original_file_path') or not self.original_file_path:
            return
            
        self.close_preview_btn.hide()
        path = self.original_file_path
        self.file_path = path
        
        try:
            self.page_data.clear()
            if path.lower().endswith('.pdf'):
                self.text_preview.hide()
                self.scroll_area.show()
                self.doc = fitz.open(path)
                
                self.status_label.setText("Restoring original PDF...")
                QApplication.processEvents()

                for i, page in enumerate(self.doc):
                    if i < 15: 
                        pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
                        img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888).copy()
                        self.page_data.append((QPixmap.fromImage(img), i))
                
                self.current_zoom = 600 
                self.refresh_pdf_view() 
            else:
                self.scroll_area.hide()
                self.text_preview.show()
                with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                    self.text_preview.setText(f.read())
                    
            self.status_label.setText("Restored original uploaded document.")
        except Exception as e:
            QMessageBox.critical(self, "Restore Error", str(e))

    def upload_file(self):
        path, _ = QFileDialog.getOpenFileName(self, "Open Document", "", "Documents (*.pdf *.txt)")
        if path:
            self.file_path = path
            self.original_file_path = path
            self.close_preview_btn.hide()
            
            self.check_upload.setChecked(False)
            for cb in self.topic_checkboxes.values():
                cb.setChecked(False)
                
            self.project_name_input.clear()
            self.intro_text_data = "" 
            self.summary_icon_btn.setObjectName("IconBtn"); self.summary_icon_btn.style().unpolish(self.summary_icon_btn); self.summary_icon_btn.style().polish(self.summary_icon_btn)
            
            self.manual_input.clear() 
            self.page_data.clear() 
            
            try:
                if path.lower().endswith('.pdf'):
                    self.text_preview.hide()
                    self.scroll_area.show()
                    self.doc = fitz.open(path)
                    
                    self.status_label.setText("Loading PDF Pages...")
                    QApplication.processEvents()

                    for i, page in enumerate(self.doc):
                        if i < 15: 
                            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
                            img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888).copy()
                            self.page_data.append((QPixmap.fromImage(img), i))
                    
                    self.current_zoom = 600 
                    self.refresh_pdf_view() 
                else:
                    self.scroll_area.hide()
                    self.text_preview.show()
                    with open(path, 'r', encoding='utf-8', errors='ignore') as f:
                        self.text_preview.setText(f.read())

                self.check_upload.setChecked(True)
                
                full_text = self.text_preview.toPlainText() if not self.doc else "".join(page.get_text() for page in self.doc)
                self.start_analysis_thread(full_text)
                
            except Exception as e:
                QMessageBox.critical(self, "Error", f"Could not load file: {str(e)}")

    def start_analysis_thread(self, text):
        self.status_label.setText("Scanning document with Python NLP...")
        self.progress_bar.show()
        self.worker = AnalysisWorker(text, self.topics)
        self.worker.finished.connect(self.handle_analysis_done)
        self.worker.error.connect(lambda e: self.status_label.setText(f"Analysis Error: {e}"))
        self.worker.start()

    def handle_analysis_done(self, results):
        self.progress_bar.hide()
        self.status_label.setText("Analysis Complete")
        
        if results["project_name"]:
            self.project_name_input.setText(results["project_name"])
            
        if results["introduction"]:
            self.intro_text_data = results["introduction"]
            self.summary_icon_btn.setObjectName("IconBtnActive"); self.summary_icon_btn.style().unpolish(self.summary_icon_btn); self.summary_icon_btn.style().polish(self.summary_icon_btn)
            
        for topic in results["found_topics"]:
            if topic in self.topic_checkboxes:
                self.topic_checkboxes[topic].setChecked(True)

    # ------------------------------------------
    # --- TEMPLATE INJECTION & EXPORT ---
    # ------------------------------------------
    def preview_pdf(self):
        try:
            import tempfile
            import webbrowser
            
            doc = QTextDocument()
            html = "<html><body style='font-family: Calibri, sans-serif; font-size: 14px;'>"
            
            p_name = self.project_name_input.text().strip() or "__________________________________________"
            intro_text = self.intro_text_data.strip() or "__________________________________________________"
            manual_text = self.manual_input.toPlainText().replace('\\n', '<br>')
            
            scope_val = getattr(self, 'scope_dropdown', None)
            scope_text = "<br>".join(scope_val.checkedItems()) if scope_val and scope_val.checkedItems() else "____"
            
            sth_val = getattr(self, 'stakeholder_dropdown', None)
            sth_text = "<br>".join(sth_val.checkedItems()) if sth_val and sth_val.checkedItems() else "Both"
            
            cert_val = getattr(self, 'cert_dropdown', None)
            cert_text = "<br>".join(cert_val.checkedItems()) if cert_val and cert_val.checkedItems() else "______________________________________________"
            
            task_val = getattr(self, 'task_dropdown', None)
            task_text = "<br>".join(task_val.checkedItems()) if task_val and task_val.checkedItems() else "______________________________________________"
            
            comm_val = getattr(self, 'comm_dropdown', None)
            comm_text = "<br>".join(comm_val.checkedItems()) if comm_val and comm_val.checkedItems() else "______________________________________________"
            
            html += f"<h1 style='text-align: center;'>TASK DIRECTIVE</h1>"
            html += f"<h3 style='text-align: center;'>PROJECT NAME: {p_name}</h3><hr>"
            
            html += "<h2>1. Introduction [Automated]</h2>"
            html += f"<p>{intro_text}</p>"
            
            html += "<h2>2. Reference [Default]</h2><p>__________________________________________________</p>"
            html += "<h2>3. Basis Of Task Directive [Default]</h2><p>__________________________________________________</p>"
            
            html += f"<h2>4. Scope Of Task Directive [Drop down menu]</h2>"
            html += f"<p>To assign the certification respectively:<br>{scope_text}<br><br>1) ____<br>2) ____<br>3) ____</p>"
            
            html += f"<h2>5. Stakeholders [Automated + Manual Notes] ({sth_text})</h2>"
            html += f"<p>The following are the major stakeholders and manual notes extracted:<br>{manual_text}</p>"
            html += "<table border='1' cellspacing='0' cellpadding='5' width='100%'><tr><th>Sl No.</th><th>Organisation</th><th>Role</th><th>Activities</th></tr><tr><td>&nbsp;</td><td></td><td></td><td></td></tr></table>"
            
            html += "<h2>6. Certification Work Breakdown [Drop down menu]</h2>"
            html += f"<p>{cert_text}</p>"
            
            html += "<h2>7. Task Allocation [Default]</h2>"
            html += f"<p>{task_text}</p>"
            
            html += "<h2>9. Communication [Default]</h2>"
            html += f"<p>{comm_text}</p>"

            html += "</body></html>"
            
            doc.setHtml(html)
            
            temp_path = os.path.join(tempfile.gettempdir(), "Task_Directive_Preview.pdf")
            printer = QPrinter(QPrinter.PrinterMode.HighResolution)
            printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
            printer.setOutputFileName(temp_path)
            
            printer.setPageMargins(QMarginsF(15.0, 15.0, 15.0, 15.0), QPageLayout.Unit.Millimeter) 
            doc.print(printer)
            
            # Load the PDF into the left preview pane exactly as requested
            try:
                self.file_path = temp_path
                self.text_preview.hide()
                self.scroll_area.show()
                self.doc = fitz.open(temp_path)
                
                self.status_label.setText("Loading PDF Preview Pages...")
                QApplication.processEvents()

                self.page_data.clear()
                for i, page in enumerate(self.doc):
                    if i < 15:
                        pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
                        img = QImage(pix.samples, pix.width, pix.height, pix.stride, QImage.Format.Format_RGB888).copy()
                        self.page_data.append((QPixmap.fromImage(img), i))
                
                self.current_zoom = 600
                self.refresh_pdf_view()
                
                self.close_preview_btn.show()
                self.status_label.setText("Preview loaded on the left pane!")
            except Exception as e:
                QMessageBox.critical(self, "Preview Load Error", str(e))
            
        except Exception as e:
            QMessageBox.critical(self, "Preview Error", f"Failed to generate preview:\\n{str(e)}")

    def export_docx(self):
        save_path, _ = QFileDialog.getSaveFileName(self, "Save Word", "Task_Directive_Final.docx", "Word (*.docx)")
        if not save_path: return
            
        try:
            def add_page_border(section):
                sectPr = section._sectPr
                pgBorders = OxmlElement('w:pgBorders')
                pgBorders.set(qn('w:offsetFrom'), 'page')
                for side in ['top', 'left', 'bottom', 'right']:
                    border = OxmlElement(f'w:{side}')
                    border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '12'); border.set(qn('w:space'), '24'); border.set(qn('w:color'), '000000')
                    pgBorders.append(border)
                sectPr.append(pgBorders)

            def remove_page_border(section):
                sectPr = section._sectPr
                for el in sectPr.findall(qn('w:pgBorders')): sectPr.remove(el)

            def add_heading(doc, text):
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True; run.font.size = Pt(12)

            doc = Document()
            style = doc.styles['Normal']
            style.font.name = 'Calibri'; style.font.size = Pt(11)

            # --- SECTION 1: COVER PAGE ---
            section1 = doc.sections[0]
            section1.top_margin = Inches(1.5); section1.bottom_margin = Inches(1.5)
            section1.left_margin = Inches(1.5); section1.right_margin = Inches(1.5)
            add_page_border(section1)

            doc.add_paragraph("File No. ________________________")
            p = doc.add_paragraph()
            run = p.add_run("\\nTASK DIRECTIVE\\n"); run.bold = True; run.font.size = Pt(20)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p = doc.add_paragraph("________ /2026"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Issue No."; table.cell(0, 1).text = "Date of Issue:"

            p_name = self.project_name_input.text().strip()
            if not p_name or p_name.lower() == "not found":
                p_name = "__________________________________________"
            doc.add_paragraph(f"\\nPROJECT NAME: {p_name}")

            p = doc.add_paragraph("\\n\\n[ LOGO HERE ]"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("ABC Organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Address of organization").alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph("Organization details").alignment = WD_ALIGN_PARAGRAPH.CENTER

            # --- SECTION 2: MAIN CONTENT ---
            section2 = doc.add_section(WD_SECTION.NEW_PAGE)
            section2.header.is_linked_to_previous = False; section2.footer.is_linked_to_previous = False
            section2.top_margin = Inches(0.75); section2.bottom_margin = Inches(0.75)
            section2.left_margin = Inches(0.75); section2.right_margin = Inches(0.75)
            remove_page_border(section2)

            add_heading(doc, "1. Introduction [Automated]")
            intro_text = self.intro_text_data.strip()
            if not intro_text: intro_text = "__________________________________________________"
            doc.add_paragraph(intro_text)

            add_heading(doc, "2. Reference [Default]")
            doc.add_paragraph("__________________________________________________")
            add_heading(doc, "3. Basis Of Task Directive [Default]")
            scope_val = getattr(self, 'scope_dropdown', None)
            scope_text = "\\n".join(scope_val.checkedItems()) if scope_val and scope_val.checkedItems() else "____"
            add_heading(doc, "4. Scope Of Task Directive [Drop down menu]")
            doc.add_paragraph(f"To assign the certification respectively:\\n{scope_text}\\n\\n1) ____\\n2) ____\\n3) ____")

            sth_val = getattr(self, 'stakeholder_dropdown', None)
            sth_text = "\\n".join(sth_val.checkedItems()) if sth_val and sth_val.checkedItems() else "Both"
            add_heading(doc, f"5. Stakeholders [Automated + Manual Notes] ({sth_text})")
            doc.add_paragraph("The following are the major stakeholders and manual notes extracted:")
            doc.add_paragraph(self.manual_input.toPlainText()) 

            table = doc.add_table(rows=4, cols=4); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No.", "Organisation", "Role", "Activities"]): table.cell(0, i).text = h

            cert_val = getattr(self, 'cert_dropdown', None)
            cert_text = "\\n".join(cert_val.checkedItems()) if cert_val and cert_val.checkedItems() else "______________________________________________"
            add_heading(doc, "6. Certification Work Breakdown [Drop down menu]")
            doc.add_paragraph(cert_text)
            
            task_val = getattr(self, 'task_dropdown', None)
            task_text = "\\n".join(task_val.checkedItems()) if task_val and task_val.checkedItems() else "______________________________________________"
            add_heading(doc, "7. Task Allocation [Default]")
            doc.add_paragraph(task_text)
            add_heading(doc, "7.1 Coordinating Directorate [Default]")
            doc.add_paragraph("______________________________________________")
            add_heading(doc, "7.2 Single Point of Contact (SPoC) [Default]")
            doc.add_paragraph("______________________________________________")

            add_heading(doc, "7.3 Certification Task Allocation")
            table = doc.add_table(rows=4, cols=4); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No.", "Certification Activity", "Certification Work Centre", "Responsible Head"]): table.cell(0, i).text = h

            add_heading(doc, "7.4 Issue of Clearance [Default]")
            doc.add_paragraph("a. ___\\nb. ___\\nc. ___\\nd. ___\\ne. ___\\nf. ___\\ng. ___")

            add_heading(doc, "8. SCRB And TARB [Default]")
            doc.add_paragraph("______________________________________________")
            
            comm_val = getattr(self, 'comm_dropdown', None)
            comm_text = "\\n".join(comm_val.checkedItems()) if comm_val and comm_val.checkedItems() else "______________________________________________"
            add_heading(doc, "9. Communication [Default]")
            doc.add_paragraph(comm_text)
            add_heading(doc, "10. Certification Progress Review [Default]")
            doc.add_paragraph("______________________________________________\\n\\n(__________)")

            add_heading(doc, "11. Distribution List")
            doc.add_paragraph("11.1 External Organization\\n1. ___\\n2. ___\\n3. ___\\n11.2 Internal Distribution")

            # --- ANNEXURES ---
            doc.add_page_break()
            add_heading(doc, "Annexure-1")
            doc.add_paragraph("Work Assignment List of LRUs [Automate + Manual]")
            table = doc.add_table(rows=4, cols=6); table.style = 'Table Grid'
            for i, h in enumerate(["Sl No", "ABC", "Abc2", "Abc3", "Abc4", "Abc5"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Integration Checks and Clearance [Manual]")
            table = doc.add_table(rows=3, cols=5); table.style = 'Table Grid'
            for i, h in enumerate(["Sl no.", "Abc1", "Abc2", "Abc3", "Abc4"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Annexure-2")
            doc.add_paragraph("Contact details of dealing officers and RDs [Manual]")
            table = doc.add_table(rows=3, cols=5); table.style = 'Table Grid'
            for i, h in enumerate(["Sl no.", "Abc1", "Abc2", "Abc3", "Abc4"]): table.cell(0, i).text = h

            doc.add_page_break()
            add_heading(doc, "Annexure-3")
            doc.add_paragraph("Product Break Down Structure [Automate]")

            doc.save(save_path)
            QMessageBox.information(self, "Success", "Task Directive Template created successfully!")

        except Exception as e:
            QMessageBox.critical(self, "Export Error", f"Failed to build template: {str(e)}")

if __name__ == "__main__":
    app = QApplication(sys.argv)

    app.setStyle(QStyleFactory.create("Fusion"))

    palette = QPalette()
    palette.setColor(QPalette.ColorRole.Base, QColor("#ffffff"))
    palette.setColor(QPalette.ColorRole.Text, QColor("#000000"))
    palette.setColor(QPalette.ColorRole.Highlight, QColor("#e2e8f0"))
    palette.setColor(QPalette.ColorRole.HighlightedText, QColor("#000000"))
    app.setPalette(palette)

    window = OfflineApp()
    window.show()

    sys.exit(app.exec())
